import os
import csv
import uuid
from typing import List, Optional
from langchain_core.documents import Document
from app.config import DATA_PATH


class NativePDFLoader:
    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        import pypdf
        docs = []
        reader = pypdf.PdfReader(self.file_path)
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                docs.append(Document(page_content=text, metadata={"source": os.path.basename(self.file_path), "page": i + 1}))
        return docs


class NativeTextLoader:
    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        base_name = os.path.basename(self.file_path)
        for enc in ("utf-8", "latin-1", "cp1252"):
            try:
                with open(self.file_path, "r", encoding=enc) as f:
                    text = f.read()
                return [Document(page_content=text, metadata={"source": base_name, "page": 1})]
            except UnicodeDecodeError:
                continue
        with open(self.file_path, "r", errors="ignore") as f:
            return [Document(page_content=f.read(), metadata={"source": base_name, "page": 1})]


class NativeCSVLoader:
    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        docs = []
        base_name = os.path.basename(self.file_path)
        with open(self.file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            headers = next(reader, None)
            for i, row in enumerate(reader):
                if headers:
                    content = "\n".join(f"{h}: {val}" for h, val in zip(headers, row))
                else:
                    content = ", ".join(row)
                if content.strip():
                    docs.append(Document(page_content=content, metadata={"source": base_name, "row": i + 1, "page": i + 1}))
        return docs


class NativeHTMLLoader:
    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        from bs4 import BeautifulSoup
        base_name = os.path.basename(self.file_path)
        with open(self.file_path, "r", encoding="utf-8", errors="ignore") as f:
            soup = BeautifulSoup(f.read(), "html.parser")
            text = soup.get_text(separator="\n")
            return [Document(page_content=text, metadata={"source": base_name, "page": 1})]


class NativeDocxLoader:
    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        import docx2txt
        base_name = os.path.basename(self.file_path)
        text = docx2txt.process(self.file_path)
        return [Document(page_content=text, metadata={"source": base_name, "page": 1})]


class NativeExcelLoader:
    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        import openpyxl
        base_name = os.path.basename(self.file_path)
        wb = openpyxl.load_workbook(self.file_path, data_only=True)
        docs = []
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            lines = []
            for row in ws.iter_rows(values_only=True):
                non_empty = [str(cell) for cell in row if cell is not None]
                if non_empty:
                    lines.append(", ".join(non_empty))
            if lines:
                docs.append(Document(page_content="\n".join(lines), metadata={"source": base_name, "sheet": sheetname, "page": 1}))
        return docs


def get_loader(file_path: str):
    file_ext = os.path.splitext(file_path)[1].lower()
    loader_mapping = {
        ".pdf": NativePDFLoader,
        ".txt": NativeTextLoader,
        ".html": NativeHTMLLoader,
        ".csv": NativeCSVLoader,
        ".xlsx": NativeExcelLoader,
        ".xls": NativeExcelLoader,
        ".docx": NativeDocxLoader,
    }
    
    if file_ext in loader_mapping:
        return loader_mapping[file_ext](file_path)
    return None


def load_documents(data_path=DATA_PATH):
    docs = []
    
    if not os.path.exists(data_path):
        print(f"Data path {data_path} does not exist.")
        return []

    for file in os.listdir(data_path):
        file_path = os.path.join(data_path, file)
        loader = get_loader(file_path)
        if loader:
            try:
                file_docs = loader.load()
                doc_id = str(uuid.uuid4())

                for i, d in enumerate(file_docs):
                    d.metadata.update({
                        "source": file,
                        "doc_id": doc_id,
                        "page": d.metadata.get("page", i + 1),
                        "category": classify_doc(file)
                    })

                docs.extend(file_docs)
            except Exception as e:
                print(f"Error loading {file}: {e}")

    return docs


def classify_doc(filename):
    name = filename.lower()
    if "india" in name or "independence" in name:
        return "history"
    return "general"
